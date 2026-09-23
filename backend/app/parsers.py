from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from .analyzer import clean_text


def parse_file(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md"}:
        return clean_text(content.decode("utf-8", errors="ignore"))
    if suffix == ".pdf":
        reader = PdfReader(BytesIO(content))
        return clean_text("\n".join(page.extract_text() or "" for page in reader.pages))
    if suffix == ".docx":
        document = Document(BytesIO(content))
        blocks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                blocks.append(" | ".join(cell.text for cell in row.cells))
        return clean_text("\n".join(blocks))
    if suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        rows: list[str] = []
        for sheet in workbook.worksheets:
            rows.append(f"Лист: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in row if value is not None]
                if values:
                    rows.append(" | ".join(values))
        return clean_text("\n".join(rows))
    raise ValueError(f"Формат {suffix or 'без расширения'} пока не поддерживается")

