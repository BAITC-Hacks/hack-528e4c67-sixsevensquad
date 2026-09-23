import json
import csv
import io
import logging
from threading import Event, Lock
from concurrent.futures import ThreadPoolExecutor
from contextlib import asynccontextmanager
from pathlib import Path

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import Response, StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from openai import (AuthenticationError, RateLimitError, APIConnectionError, APIStatusError, APITimeoutError,
                    BadRequestError, LengthFinishReasonError, ContentFilterFinishReasonError)
from pymongo.errors import PyMongoError

from .config import ROOT, Settings
from .schemas import AnalysisRequest, Document, Review
from .parsers import parse_document, MAX_BYTES
from .pipeline import analyze, build_report
from .provider import OpenAIProvider, AnalysisCancelled
from .baseline import BaselineProvider
from .storage import Store, StorageUnavailable
from .memory_storage import MemoryStore
from .local_storage import LocalStore
from .storage import now

def create_app(settings=None, provider_factory=None, store=None):
    settings = settings or Settings()
    store = store if store is not None else (LocalStore(settings.kontur_data_dir) if settings.kontur_storage == "local" else MemoryStore() if settings.kontur_storage == "memory" else Store(settings))
    executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="kontur")
    job_lock = Lock()
    cancellations = {}
    @asynccontextmanager
    async def lifespan(app):
        try:
            if store.configured:
                try:
                    store.initialize()
                    store.recover()
                except (PyMongoError, ValueError):
                    raise RuntimeError("Хранилище недоступно: проверьте локальные файлы или настройки MongoDB") from None
            yield
        finally:
            for event in list(cancellations.values()):
                event.set()
            executor.shutdown(wait=True, cancel_futures=False)
            store.close()
    app = FastAPI(title="Контур API", version="0.1.0", lifespan=lifespan)
    app.add_middleware(CORSMiddleware, allow_origins=["http://localhost:3000", "http://127.0.0.1:3000"],
                       allow_methods=["*"], allow_headers=["*"])
    app.state.store = store

    @app.exception_handler(StorageUnavailable)
    @app.exception_handler(PyMongoError)
    async def database_unavailable(request, exc):
        from fastapi.responses import JSONResponse
        return JSONResponse(status_code=503, content={"detail":"MongoDB недоступна. Проверьте MONGODB_URI, доступ к кластеру и перезапустите сервер"})

    @app.exception_handler(KeyError)
    async def missing(request, exc):
        from fastapi.responses import JSONResponse
        return JSONResponse(status_code=404, content={"detail":"Проект или вывод не найден"})

    @app.get("/api/health")
    def health():
        database = "not_configured"
        if store.configured:
            try:
                store.ping()
                database = "local" if isinstance(store, LocalStore) else "memory" if isinstance(store, MemoryStore) else "connected"
            except (PyMongoError, ValueError):
                database = "unavailable"
        return {"status":"ok" if database in {"connected", "memory", "local"} else "degraded", "database":database,
            "openai_configured":bool(settings.openai_api_key.get_secret_value()), "model":settings.openai_model,
            "formats":["txt","docx","pdf","xlsx"], "max_file_mb":10,
            "budget_usd":settings.kontur_budget_usd, "max_calls":settings.kontur_max_calls,
            "analysis_seconds":settings.kontur_analysis_seconds}

    @app.get("/api/projects")
    def projects():
        return store.list()

    def checked_create(documents):
        if sum(len(c.text) for d in documents for c in d.clauses) > settings.kontur_max_chars:
            raise HTTPException(413, "Суммарный текст превышает лимит. Разделите комплект")
        p = store.create(documents)
        return {"id":p["id"], "status":p["status"]}

    @app.post("/api/projects", status_code=201)
    async def upload(before: list[UploadFile] = File(...), after: list[UploadFile] = File(...)):
        if len(before) + len(after) > 12:
            raise HTTPException(413, "Максимум 12 файлов в комплекте")
        documents = []
        try:
            for side, files in [("before", before), ("after", after)]:
                for i, file in enumerate(files):
                    data = await file.read(MAX_BYTES + 1)
                    documents.append(parse_document(file.filename or "document", data, side, f"{side}-{i+1}"))
        except Exception as exc:
            if isinstance(exc, ValueError):
                raise HTTPException(422, str(exc)) from None
            raise HTTPException(422, "Не удалось прочитать файл. Проверьте формат и целостность документа") from None
        finally:
            for f in before + after:
                await f.close()
        return checked_create(documents)

    @app.post("/api/projects/sample", status_code=201)
    def sample(kind: str = "organizer"):
        documents = []
        if kind not in {"organizer", "control"}:
            raise HTTPException(422, "Пример: organizer или control")
        for side in ("before", "after"):
            path = ROOT / "data" / "samples" / (f"organizer-{side}.txt" if kind == "organizer" else f"control-{side}.txt")
            if not path.is_file():
                raise HTTPException(404, "Файлы примера отсутствуют")
            documents.append(parse_document(path.name, path.read_bytes(), side, f"{side}-1"))
        return checked_create(documents)

    @app.get("/api/projects/{pid}")
    def project(pid: str):
        p = store.get(pid)
        p.pop("_cache", None)
        if p["result"]:
            p["result"]["report"] = build_report(p["result"], p["reviews"])
        return p

    @app.get("/api/projects/{pid}/status")
    def project_status(pid: str):
        return store.status(pid)

    @app.post("/api/projects/{pid}/cancel")
    def cancel(pid: str):
        with job_lock:
            p = store.get(pid)
            if p["status"] != "running" or pid not in cancellations:
                raise HTTPException(409, "Этот анализ сейчас не выполняется")
            cancellations[pid].set()
            store.update(pid, stage="Остановка: ожидаем завершения текущего запроса")
        return {"status":"cancelling"}

    def run(pid, mode):
        provider = None
        try:
            event = cancellations[pid]
            provider = provider_factory(mode) if provider_factory else OpenAIProvider(settings,
                cache_get=lambda key:store.cache_get(pid, key),
                cache_put=lambda key,value:store.cache_put(pid, key,value),
                metrics=store.get(pid).get("metrics"),
                on_metrics=lambda metrics:store.update(pid, metrics=metrics),
                cancelled=event.is_set) if mode == "openai" else BaselineProvider()
            documents = [Document.model_validate(d) for d in store.get(pid)["documents"]]
            def progress(stage):
                if event.is_set():
                    raise AnalysisCancelled("Анализ остановлен; завершённые запросы сохранены")
                store.update(pid, stage=stage)
            result = analyze(documents, provider, mode, progress)
            if event.is_set():
                raise AnalysisCancelled("Анализ остановлен; завершённые запросы сохранены")
            store.update(pid, status="completed", stage="Анализ завершён с замечаниями" if result.get("warnings") else "Анализ завершён", result=result)
        except AnalysisCancelled as exc:
            store.update(pid, status="cancelled", error=None, stage=str(exc))
        except AuthenticationError:
            store.update(pid, status="failed", error="OpenAI отклонил ключ. Проверьте OPENAI_API_KEY в .env", stage="Ошибка авторизации")
        except RateLimitError:
            store.update(pid, status="failed", error="OpenAI: квота или лимит запросов. Проверьте биллинг и повторите позднее", stage="Лимит OpenAI")
        except APITimeoutError:
            store.update(pid, status="failed", error="OpenAI не ответил за отведённое время. Готовые ответы сохранены; можно продолжить анализ", stage="Время ожидания OpenAI истекло")
        except APIConnectionError:
            store.update(pid, status="failed", error="Нет соединения с OpenAI. Проверьте сеть", stage="Ошибка сети")
        except BadRequestError:
            store.update(pid, status="failed", error="OpenAI не принял запрос. Проверьте выбранную модель, доступ к ней и размер комплекта", stage="Ошибка запроса OpenAI")
        except APIStatusError:
            store.update(pid, status="failed", error="OpenAI отклонил запрос. Проверьте доступ к OPENAI_MODEL и повторите", stage="Ошибка OpenAI")
        except LengthFinishReasonError:
            store.update(pid, status="failed", error="Ответ модели превысил лимит. Разделите комплект на меньшие части", stage="Лимит ответа ИИ")
        except ContentFilterFinishReasonError:
            store.update(pid, status="failed", error="Модель не вернула результат из-за ограничения обработки", stage="Ограничение ИИ")
        except ValueError as exc:
            store.update(pid, status="failed", error=str(exc)[:1000], stage="Проверка анализа не пройдена")
        except Exception:
            logging.getLogger(__name__).exception("Analysis failed for project %s", pid)
            store.update(pid, status="failed", error="Внутренняя ошибка обработки. Результат не опубликован", stage="Ошибка")
        finally:
            try:
                store.update(pid, finished=now())
            finally:
                with job_lock:
                    cancellations.pop(pid, None)

    @app.post("/api/projects/{pid}/analyze", status_code=202)
    def start(pid: str, request: AnalysisRequest):
        if request.mode == "openai" and not settings.openai_api_key.get_secret_value() and not provider_factory:
            raise HTTPException(503, "Добавьте OPENAI_API_KEY в .env и перезапустите сервер. Ключ не отправляется в браузер")
        with job_lock:
            if cancellations:
                raise HTTPException(409, "Уже выполняется анализ. Дождитесь завершения или остановите его в истории")
            try:
                p = store.start(pid, request.mode)
            except ValueError as exc:
                raise HTTPException(409, str(exc)) from None
            cancellations[pid] = Event()
            executor.submit(run, pid, request.mode)
        return {"id":pid, "status":p["status"]}

    @app.put("/api/projects/{pid}/findings/{fid}/review")
    def review(pid: str, fid: str, request: Review):
        return store.review(pid, fid, request.model_dump())

    @app.get("/api/projects/{pid}/export")
    def export(pid: str, format: str = "md"):
        p = project(pid)
        if not p["result"]:
            raise HTTPException(409, "Анализ ещё не завершён")
        if format == "json":
            body, media = json.dumps(p, ensure_ascii=False, indent=2), "application/json"
        elif format == "md":
            body, media = build_report(p["result"], p["reviews"]), "text/markdown"
        elif format == "csv":
            output = io.StringIO()
            writer = csv.writer(output, delimiter=";")
            writer.writerow(["ID", "Статус", "Функция", "Пояснение", "Рекомендация", "Источники", "Решение аналитика"])
            for finding in p["result"]["findings"]:
                sources = " | ".join(f"{e['document']}, {e['locator']}: {e['quote']}" for e in finding["evidence"])
                row = [finding["id"], finding["label"], finding["title"], finding["explanation"],
                                 finding["recommendation"], sources,
                                 p["reviews"].get(finding["id"], {}).get("decision", "pending")]
                writer.writerow(["'"+str(value) if str(value).lstrip().startswith(("=", "+", "-", "@", "\t", "\r")) else value for value in row])
            body, media = "\ufeff" + output.getvalue(), "text/csv; charset=utf-8"
        elif format == "docx":
            from docx import Document as WordDocument
            document = WordDocument()
            for line in build_report(p["result"], p["reviews"]).splitlines():
                if line.startswith("# "):
                    document.add_heading(line[2:], level=1)
                elif line.startswith("## "):
                    document.add_heading(line[3:], level=2)
                elif line.startswith("- "):
                    document.add_paragraph(line[2:], style="List Bullet")
                elif line:
                    document.add_paragraph(line.lstrip("> "))
            output = io.BytesIO()
            document.save(output)
            output.seek(0)
            return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                     headers={"Content-Disposition": f'attachment; filename="kontur-{pid}.docx"'})
        else:
            raise HTTPException(422, "Формат экспорта: md, json, csv или docx")
        return Response(body, media_type=media, headers={"Content-Disposition":f'attachment; filename="kontur-{pid}.{format}"'})

    @app.get("/", include_in_schema=False)
    def index():
        return {"app": "Контур API", "frontend": "http://localhost:3000", "docs": "/docs"}
    return app

app = create_app()
