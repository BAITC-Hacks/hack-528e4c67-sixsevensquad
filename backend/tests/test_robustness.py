import io
import json
import threading
import time
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from backend.app.baseline import BaselineProvider
from backend.app.config import Settings
from backend.app.evidence import EvidenceError, exact_source_quote
from backend.app.local_storage import LocalStore
from backend.app.main import create_app
from backend.app.memory_storage import MemoryStore
from backend.app.parsers import parse_document
from backend.app.pipeline import analyze, extraction_context
from backend.app.provider import OpenAIProvider, AnalysisCancelled
from backend.app.schemas import Matches, UnitChanges, Match
from backend.tests.test_mvp import fixture_documents


@pytest.mark.parametrize("variant", ["missing", "duplicate", "unknown", "wrong_kind"])
def test_bad_unit_response_preserves_result_with_explicit_uncertainty(variant):
    class Provider(BaselineProvider):
        def units(self, before, after):
            result = super().units(before, after)
            if variant == "missing": result.changes = []
            if variant == "duplicate": result.changes.append(result.changes[0])
            if variant == "unknown": result.changes[0].before_names = ["Нет такого подразделения"]
            if variant == "wrong_kind": result.changes[0].status = "created"
            return result
    result = analyze(fixture_documents("control"), Provider(), "openai")
    assert result["findings"] and result["warnings"]
    for side in ("before", "after"):
        names = [n for c in result["unit_changes"] for n in c[side+"_names"]]
        assert len(names) == len(set(names))
        assert set(names) == {u["name"] for u in result["units"][side]}


@pytest.mark.parametrize("variant", ["missing", "duplicate", "unknown_after"])
def test_bad_matches_are_uncertain_not_losses(variant):
    class Provider(BaselineProvider):
        def match(self, before, after):
            result = super().match(before, after)
            if variant == "missing": result.matches = []
            if variant == "duplicate": result.matches.append(result.matches[0])
            if variant == "unknown_after": result.matches[0].after_ids = ["fabricated"]
            return result
    result = analyze(fixture_documents("control"), Provider(), "openai")
    assert result["counts"]["uncertain"] >= 1
    matched = [i for f in result["findings"] for i in f["before_ids"]]
    assert len(matched) == len(set(matched)) == result["coverage"]["before_functions"]


def test_strict_quotes_allow_whitespace_only_not_invention():
    assert exact_source_quote("Ведёт   реестр договоров.", "Ведёт реестр договоров") == "Ведёт   реестр договоров"
    with pytest.raises(EvidenceError): exact_source_quote("Ведёт реестр договоров.", "Утверждает бюджет отдела")


def test_openai_units_are_partitioned_by_kind(monkeypatch):
    payloads = []
    def parse(**kwargs):
        payload = json.loads(kwargs["input"][1]["content"])
        payloads.append(payload)
        kinds = {u["kind"] for u in payload["before"] + payload["after"]}
        assert len(kinds) == 1
        return SimpleNamespace(status="completed",
            output_parsed=BaselineProvider().units(payload["before"], payload["after"]),
            usage=SimpleNamespace(input_tokens=1, output_tokens=1))
    monkeypatch.setattr("backend.app.provider.OpenAI", lambda **kw: SimpleNamespace(responses=SimpleNamespace(parse=parse)))
    provider = OpenAIProvider(Settings(_env_file=None, openai_api_key="test"))
    before = [{"name":"Отдел аудита", "kind":"department"}, {"name":"Начальник отдела", "kind":"role"}]
    after = before + [{"name":"Компания", "kind":"organization"}]
    result = provider.units(before, after)
    assert len(payloads) == 3
    for side, registry in (("before", before), ("after", after)):
        assert {n for c in result.changes for n in getattr(c, side + "_names")} == {u["name"] for u in registry}
    assert provider.unit_changes == [c.model_dump() for c in result.changes]


def test_risk_prompt_retains_original_independence_condition(monkeypatch):
    from backend.app.schemas import Risks
    captured = {}
    def parse(**kwargs):
        captured.update(json.loads(kwargs["input"][1]["content"]))
        return SimpleNamespace(status="completed", output_parsed=Risks(risks=[]),
                               usage=SimpleNamespace(input_tokens=1, output_tokens=1))
    monkeypatch.setattr("backend.app.provider.OpenAI", lambda **kw: SimpleNamespace(responses=SimpleNamespace(parse=parse)))
    provider = OpenAIProvider(Settings(_env_file=None, openai_api_key="test"))
    f = {"id":"f1", "owner":"Отдел", "action":"проверяет", "object":"решения", "scope":"не указана",
         "evidence":[{"clause_id":"after:c1", "quote":"Независимо проверяет собственные решения."}]}
    provider.risks([f], [f])
    assert captured["focus"][0]["evidence"] == f["evidence"]
    assert "evidence" not in captured["catalog"][0]


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-16"])
def test_txt_encodings_and_long_lines(encoding):
    text = "1. Логистика: " + "Ведёт реестр грузов. "*500
    doc = parse_document("ДРУГОЙ.TXT", text.encode(encoding), "before", "d")
    assert len(doc.clauses) > 1
    assert max(len(c.text) for c in doc.clauses) <= 6000
    assert " ".join(c.text for c in doc.clauses) == text.strip()


@pytest.mark.parametrize("name,data", [("empty.txt", b""), ("fake.docx", b"not zip"), ("old.doc", b"a"*50), ("binary.txt", b"\x00"*40)])
def test_invalid_uploads_are_explained(name, data):
    with TestClient(create_app(Settings(_env_file=None), store=MemoryStore())) as client:
        response = client.post("/api/projects", files=[("before",(name,data)),("after",("good.txt", b"Some responsibility description for a department."))])
        assert response.status_code == 422
        assert response.json()["detail"]


def test_pdf_with_text_and_mixed_scan():
    from reportlab.pdfgen.canvas import Canvas
    buffer = io.BytesIO()
    canvas = Canvas(buffer)
    canvas.drawString(50,700,"1.1. Logistics department maintains the shipping register.")
    canvas.save()
    doc = parse_document("logistics.pdf",buffer.getvalue(),"before","pdf")
    assert "shipping" in doc.clauses[0].text and "стр. 1" in doc.clauses[0].locator
    from pypdf import PdfReader, PdfWriter
    output = io.BytesIO(); writer = PdfWriter()
    writer.add_page(PdfReader(io.BytesIO(buffer.getvalue())).pages[0]); writer.add_blank_page(width=300,height=300); writer.write(output)
    with pytest.raises(ValueError, match="страница 2"):
        parse_document("mixed.pdf",output.getvalue(),"before","pdf")


def test_excel_excessive_columns_are_rejected():
    from openpyxl import Workbook
    book = Workbook(); book.active.cell(1,201,"test")
    stream = io.BytesIO(); book.save(stream)
    with pytest.raises(ValueError, match="столбцов"):
        parse_document("wide.xlsx",stream.getvalue(),"before","xlsx")


def test_local_history_cache_reviews_survive_restart(tmp_path):
    store = LocalStore(tmp_path); store.initialize()
    documents = fixture_documents("control")
    p = store.create(documents); pid = p["id"]
    store.cache_put(pid,"key",{"answer":"cached"})
    store.start(pid,"baseline")
    store.update(pid,status="completed",result=analyze(documents,BaselineProvider(),"baseline"))
    store.review(pid,"finding-1",{"decision":"confirmed","comment":"Verified"})
    second = LocalStore(tmp_path); second.initialize(); second.recover()
    assert second.cache_get(pid,"key") == {"answer":"cached"}
    assert second.list()[0]["reviewed"] == 1
    assert "control-before.txt" in second.list()[0]["title"]
    assert second.get(pid)["reviews"]["finding-1"]["comment"] == "Verified"
    assert (tmp_path/f"{pid}.json").stat().st_mode & 0o777 == 0o600


def test_local_interrupted_jobs_are_recoverable(tmp_path):
    store = LocalStore(tmp_path); store.initialize()
    p = store.create(fixture_documents("control")); store.start(p["id"],"openai")
    store.cache_put(p["id"],"key",{"answer":1})
    second = LocalStore(tmp_path); second.initialize(); second.recover()
    assert second.get(p["id"])["status"] == "failed"
    assert second.cache_get(p["id"],"key") == {"answer":1}


def fake_provider(monkeypatch, **overrides):
    calls = []
    def parse(**kwargs):
        calls.append(kwargs)
        return SimpleNamespace(status="completed", output_parsed=Matches(matches=[]), usage=SimpleNamespace(input_tokens=100,output_tokens=20))
    monkeypatch.setattr("backend.app.provider.OpenAI",lambda **kw:SimpleNamespace(responses=SimpleNamespace(parse=parse)))
    return OpenAIProvider(Settings(_env_file=None,openai_api_key="test",**overrides)), calls


def test_budget_prevents_transport_and_retry_does_not_reset_spend(monkeypatch):
    provider,calls = fake_provider(monkeypatch,kontur_budget_usd=.000001)
    with pytest.raises(ValueError,match="бюджет"): provider.ask(Matches,"test",{})
    assert calls == []
    provider,calls = fake_provider(monkeypatch,kontur_max_calls=1)
    provider.ask(Matches,"first",{})
    with pytest.raises(ValueError,match="лимит запросов"): provider.ask(Matches,"second",{})
    assert len(calls) == 1
    again = OpenAIProvider(Settings(_env_file=None,openai_api_key="test",kontur_max_calls=1),metrics=provider.snapshot())
    with pytest.raises(ValueError,match="лимит запросов"): again.ask(Matches,"third",{})


def test_cache_hits_are_free_and_still_validate_schema(monkeypatch):
    provider,calls = fake_provider(monkeypatch)
    cache = {}; provider.cache_get = cache.get; provider.cache_put = lambda k,v:cache.update({k:v})
    provider.ask(Matches,"test",{}); provider.ask(Matches,"test",{})
    assert len(calls) == 1 and provider.cache_hits == 1
    assert provider.usage["input_tokens"] == 100
    assert provider.reserved == pytest.approx(0)


def test_deadline_and_cancel_block_new_calls(monkeypatch):
    provider,calls = fake_provider(monkeypatch)
    provider.cancelled = lambda:True
    with pytest.raises(AnalysisCancelled): provider.ask(Matches,"test",{})
    provider.cancelled = lambda:False; provider.started -= 700
    with pytest.raises(ValueError,match="времени"): provider.ask(Matches,"test",{})
    assert not calls


def test_api_single_active_analysis_cancellation_and_private_cache():
    entered, release = threading.Event(), threading.Event()
    class Slow(BaselineProvider):
        def extract(self, clauses, context):
            entered.set(); release.wait(3)
            return super().extract(clauses,context)
    store = MemoryStore()
    with TestClient(create_app(Settings(_env_file=None),provider_factory=lambda mode:Slow(),store=store)) as client:
        first = client.post("/api/projects/sample?kind=control").json()["id"]
        second = client.post("/api/projects/sample?kind=control").json()["id"]
        store.cache_put(first,"private",{"answer":1})
        assert "_cache" not in client.get(f"/api/projects/{first}").json()
        assert client.post(f"/api/projects/{first}/analyze",json={"mode":"baseline"}).status_code == 202
        assert entered.wait(2)
        assert client.post(f"/api/projects/{second}/analyze",json={"mode":"baseline"}).status_code == 409
        assert client.post(f"/api/projects/{first}/cancel").status_code == 200
        release.set()
        for _ in range(100):
            p = client.get(f"/api/projects/{first}/status").json()
            if p["status"] != "running":break
            time.sleep(.02)
        assert p["status"] == "cancelled"
        assert "control-before.txt" in p["title"]


def test_mongo_history_and_cache_contract():
    import mongomock
    from mongomock.gridfs import enable_gridfs_integration
    from backend.app.storage import Store
    enable_gridfs_integration()
    store = Store(Settings(_env_file=None),client=mongomock.MongoClient())
    store.initialize(); p = store.create(fixture_documents("control")); pid = p["id"]
    store.start(pid,"baseline"); store.cache_put(pid,"key",{"answer":1}); store.cache_put(pid,"key",{"answer":2})
    store.update(pid,status="completed",result=analyze(fixture_documents("control"),BaselineProvider(),"baseline"))
    store.review(pid,"finding-1",{"decision":"confirmed","comment":"ok"})
    assert store.list()[0]["title"] == "control-before.txt → control-after.txt"
    assert store.list()[0]["reviewed"] == 1
    assert store.cache_get(pid,"key") == {"answer":2}
    assert store.get(pid)["result"]["findings"]


def test_incomplete_extraction_splits_without_losing_fragments():
    from backend.app.provider import IncompleteResponse
    class Provider(BaselineProvider):
        sizes = []
        def extract(self, clauses, context):
            self.sizes.append(len(clauses))
            if len(clauses) > 5:
                raise IncompleteResponse("too long")
            return super().extract(clauses, context)
    provider = Provider()
    result = analyze(fixture_documents("control"), provider, "openai")
    assert result["coverage"]["clauses_processed"] == 16
    assert result["coverage"]["before_functions"] >= 3
    assert any(n > 5 for n in provider.sizes) and any(n <= 5 for n in provider.sizes)


def test_parallel_budget_reservation_is_atomic(monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    provider, calls = fake_provider(monkeypatch,kontur_budget_usd=.014)
    gate = threading.Event()
    def slow(**kwargs):
        calls.append(kwargs); gate.wait(1)
        return SimpleNamespace(status="completed",output_parsed=Matches(matches=[]),usage=SimpleNamespace(input_tokens=1,output_tokens=1))
    provider.client.responses.parse = slow
    def ask(i):
        try:provider.ask(Matches,str(i),{})
        except ValueError:return "blocked"
        return "done"
    with ThreadPoolExecutor(max_workers=2) as executor:
        jobs = [executor.submit(ask,i) for i in range(2)]
        time.sleep(.05); gate.set()
        outcomes = [j.result() for j in jobs]
    assert sorted(outcomes) == ["blocked","done"]
    assert len(calls) == 1


def test_many_uploads_rejected_without_creating_project():
    with TestClient(create_app(Settings(_env_file=None),store=MemoryStore())) as client:
        files = [("before",(f"{i}.txt",b"Department reviews shipping documents.")) for i in range(12)]
        files.append(("after",("after.txt",b"Department reviews shipping documents.")))
        assert client.post("/api/projects",files=files).status_code == 413
        assert client.get("/api/projects").json() == []


def test_csv_formula_is_not_executable():
    store = MemoryStore(); docs = fixture_documents("control"); p = store.create(docs)
    result = analyze(docs,BaselineProvider(),"baseline")
    result["findings"][0]["title"] = '=HYPERLINK("https://example.invalid")'
    store.update(p["id"],status="completed",result=result)
    with TestClient(create_app(Settings(_env_file=None),store=store)) as client:
        response = client.get(f'/api/projects/{p["id"]}/export?format=csv')
        assert "'=HYPERLINK" in response.text


def test_unique_verbatim_quote_can_correct_known_wrong_clause_id():
    from backend.app.schemas import Evidence, Extraction, Unit
    from backend.app.evidence import anchor_extraction
    doc = parse_document("a.txt","Отдел логистики ведёт реестр грузов.\nОтдел кадров ведёт реестр сотрудников.".encode(),"before","d")
    result = Extraction(units=[Unit(name="Логистика",kind="department",evidence=[Evidence(clause_id="d:c2",quote="Отдел логистики ведёт реестр грузов.")])],functions=[])
    anchor_extraction(result,doc.clauses)
    assert result.units[0].evidence[0].clause_id == "d:c1"


def test_failed_quote_repair_excludes_record_and_marks_partial():
    class Provider(BaselineProvider):
        def extract(self, clauses, context):
            result = super().extract(clauses,context)
            if clauses[0].side == "after" and result.functions:
                result.functions[0].evidence[0].quote = "Такой цитаты в документах вообще нет"
            return result
        def repair_extraction(self, clauses, context, previous, error):
            return previous
    result = analyze(fixture_documents("control"),Provider(),"openai")
    assert result["coverage"]["incomplete_sides"] == ["after"]
    assert result["warnings"] and result["counts"]["uncertain"] >= 1
    assert result["counts"].get("not_found",0) == 0
    assert not any("Такой цитаты" in e["quote"] for f in result["findings"] for e in f["evidence"])


def test_partial_extraction_does_not_claim_disappeared_units():
    from backend.app.schemas import UnitChange
    class Provider(BaselineProvider):
        def extract(self, clauses, context):
            result = super().extract(clauses, context)
            if clauses[0].side == "after" and result.functions:
                result.functions[0].evidence[0].quote = "Неподтверждённая цитата об ответственности"
            return result
        def repair_extraction(self, clauses, context, previous, error):
            from backend.app.provider import IncompleteResponse
            raise IncompleteResponse("repair truncated")
        def units(self, before, after):
            return UnitChanges(changes=[
                UnitChange(before_names=[u["name"]], after_names=[], status="not_found", explanation="нет соответствия")
                for u in before])
    result = analyze(fixture_documents("control"), Provider(), "openai")
    assert result["coverage"]["incomplete_sides"] == ["after"]
    assert result["unit_changes"]
    assert all(c["status"] == "uncertain" for c in result["unit_changes"])
    assert any("обрезан" in w for w in result["warnings"])


def test_missing_structure_registry_is_explicit_not_silent_success():
    class Provider(BaselineProvider):
        def extract(self, clauses, context):
            result = super().extract(clauses, context)
            result.units = []
            return result
    result = analyze(fixture_documents("control"), Provider(), "openai")
    assert result["findings"]
    assert sum("Не извлечён реестр" in w for w in result["warnings"]) == 2


@pytest.mark.parametrize("extension", ["docx", "xlsx"])
def test_uploaded_office_pair_runs_through_api_and_preserves_sources(extension):
    def office_file(document):
        stream = io.BytesIO()
        if extension == "docx":
            from docx import Document
            book = Document()
            for clause in document.clauses:
                book.add_paragraph(clause.text)
        else:
            from openpyxl import Workbook
            book = Workbook()
            for clause in document.clauses:
                book.active.append([clause.text])
        book.save(stream)
        return stream.getvalue()
    documents = fixture_documents("control")
    with TestClient(create_app(Settings(_env_file=None), store=MemoryStore())) as client:
        files = [(d.side, (f"independent-{d.side}.{extension.upper()}", office_file(d))) for d in documents]
        response = client.post("/api/projects", files=files)
        assert response.status_code == 201, response.text
        pid = response.json()["id"]
        assert client.post(f"/api/projects/{pid}/analyze", json={"mode": "baseline"}).status_code == 202
        for _ in range(200):
            project = client.get(f"/api/projects/{pid}").json()
            if project["status"] != "running":
                break
            time.sleep(.02)
        assert project["status"] == "completed", project.get("error")
        assert project["result"]["counts"].get("duplication", 0) > 0
        for finding in project["result"]["findings"]:
            assert finding["evidence"]
            assert all(e["quote"] in e["text"] for e in finding["evidence"])
            assert all(e["document"].endswith(extension.upper()) for e in finding["evidence"])
