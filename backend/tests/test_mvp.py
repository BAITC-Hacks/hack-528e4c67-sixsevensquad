import io
import time
from pathlib import Path
import pytest

from fastapi.testclient import TestClient

from backend.app.baseline import BaselineProvider
from backend.app.evidence import EvidenceError, anchor_extraction, validate_evidence
from backend.app.main import create_app
from backend.app.memory_storage import MemoryStore
from backend.app.parsers import parse_document
from backend.app.pipeline import analyze, retain_target_supported
from backend.app.schemas import Evidence, Extraction, Function, Match, Matches, Unit

ROOT = Path(__file__).resolve().parents[2]


def test_pdf_words_are_reassembled_without_losing_source_text():
    from backend.app.parsers import pdf_blocks
    from backend.app.evidence import normalize

    raw = "ПОЛОЖЕНИЕ\n1.1. Отдел ведёт\n \nреестр\n \nстраховых\n \nдоговоров.  1.2. Проверяет расходы.\nа. Контролирует\nдоговоры."
    blocks = pdf_blocks(raw, 3)
    assert len(blocks) == 4
    assert blocks[1][1] == "1.1. Отдел ведёт реестр страховых договоров."
    assert all("стр. 3" in locator for locator, _ in blocks)
    assert normalize(" ".join(text for _, text in blocks)) == normalize(raw)


def test_paraphrased_citation_is_anchored_to_exact_source():
    class ParaphraseProvider(BaselineProvider):
        extracts = 0

        def extract(self, clauses, context):
            result = super().extract(clauses, context)
            self.extracts += 1
            if self.extracts == 1 and result.functions:
                result.functions[0].evidence[0].quote = "Пересказ, которого буквально нет в документе"
            return result

    provider = ParaphraseProvider()
    stages = []
    result = analyze(fixture_documents("control"), provider, "openai", stages.append)
    assert result["findings"]
    assert all(e["quote"] in e["text"] for f in result["findings"] for e in f["evidence"])
    assert not any("Уточнение цитат" in stage for stage in stages)


def test_unknown_clause_id_is_still_rejected():
    result = BaselineProvider().extract(fixture_documents("control")[0].clauses, {})
    result.functions[0].evidence[0].clause_id = "before:c999999"
    with pytest.raises(EvidenceError, match="Несуществующий источник"):
        anchor_extraction(result, fixture_documents("control")[0].clauses)


def test_context_only_records_are_ignored_in_current_batch():
    result = Extraction(
        units=[Unit(name="Отдел из контекста", kind="department",
                    evidence=[Evidence(clause_id="context:c1", quote="Отдел из контекста")])],
        functions=[
            Function(owner="Отдел из контекста", action="ведёт", object="реестр", scope="договоры",
                     evidence=[Evidence(clause_id="context:c1", quote="Отдел из контекста")]),
            Function(owner="Текущий отдел", action="проверяет", object="отчёты", scope="филиалы",
                     evidence=[Evidence(clause_id="target:c1", quote="Проверяет отчёты филиалов")]),
        ],
    )
    retain_target_supported(result, {"target:c1"})
    assert result.units == []
    assert [item.owner for item in result.functions] == ["Текущий отдел"]


def fixture_documents(prefix):
    return [
        parse_document(f"{prefix}-{side}.txt", (ROOT / "data" / "samples" / f"{prefix}-{side}.txt").read_bytes(), side, side)
        for side in ("before", "after")
    ]


def test_control_set_covers_required_risks_and_sources():
    result = analyze(fixture_documents("control"), BaselineProvider(), "baseline")
    assert result["counts"]["not_found"] >= 1
    assert result["counts"]["duplication"] >= 1
    assert result["counts"]["conflict"] >= 1
    assert result["counts"]["transferred"] >= 1
    assert any(change["status"] == "reorganized" and
               change["before_names"] == ["Отдел финансового контроля"] and
               set(change["after_names"]) == {"Отдел операционного аудита", "Отдел риск-контроля"}
               for change in result["unit_changes"])
    assert all(f["evidence"] and all(e["quote"] in e["text"] for e in f["evidence"]) for f in result["findings"])


def test_organizer_structure_is_not_polluted_by_incidental_mentions():
    result = analyze(fixture_documents("organizer"), BaselineProvider(), "baseline")
    changes = result["unit_changes"]
    assert len(changes) == 5
    assert sum(c["status"] == "preserved" for c in changes) == 2
    assert sum(c["status"] == "created" for c in changes) == 2
    assert any("ДИТААД" in c["after_names"][0] for c in changes if c["status"] == "created")
    assert any("ДОА" in c["after_names"][0] for c in changes if c["status"] == "created")


def test_unverified_quote_is_rejected():
    doc = fixture_documents("control")[0]
    try:
        validate_evidence([Evidence(clause_id=doc.clauses[0].id, quote="Выдуманное правило")],
                          {c.id: c for c in doc.clauses})
    except EvidenceError:
        pass
    else:
        raise AssertionError("Неподтверждённая цитата прошла проверку")


def test_api_upload_review_and_exports():
    app = create_app(store=MemoryStore())
    with TestClient(app) as client:
        response = client.post("/api/projects/sample?kind=control")
        assert response.status_code == 201
        pid = response.json()["id"]
        assert client.post(f"/api/projects/{pid}/analyze", json={"mode": "baseline"}).status_code == 202
        for _ in range(100):
            project = client.get(f"/api/projects/{pid}").json()
            if project["status"] != "running":
                break
            time.sleep(.02)
        assert project["status"] == "completed", project.get("error")
        finding_id = project["result"]["findings"][0]["id"]
        review = client.put(f"/api/projects/{pid}/findings/{finding_id}/review",
                            json={"decision": "confirmed", "comment": "Проверено на исходнике"})
        assert review.status_code == 200
        assert "Проверено на исходнике" in client.get(f"/api/projects/{pid}/export?format=md").text
        assert client.get(f"/api/projects/{pid}/export?format=json").status_code == 200
        assert client.get(f"/api/projects/{pid}/export?format=csv").content.startswith(b"\xef\xbb\xbf")
        docx = client.get(f"/api/projects/{pid}/export?format=docx")
        assert docx.status_code == 200 and docx.content.startswith(b"PK")


def test_upload_parse_word_excel_and_reject_scan():
    from docx import Document
    from openpyxl import Workbook
    from pypdf import PdfWriter

    word = Document()
    word.add_paragraph("1.1. Отдел контроля:")
    word.add_table(rows=1, cols=1).cell(0, 0).text = "Проверяет расходы филиалов."
    stream = io.BytesIO()
    word.save(stream)
    parsed = parse_document("case.docx", stream.getvalue(), "before", "d1")
    assert "строка таблицы" in parsed.clauses[1].locator

    book = Workbook()
    book.active.append(["Владелец", "Функция"])
    book.active.append(["Отдел контроля", "Проверяет расходы"])
    stream = io.BytesIO()
    book.save(stream)
    parsed = parse_document("case.xlsx", stream.getvalue(), "after", "d2")
    assert "строка 2" in parsed.clauses[1].locator

    pdf = PdfWriter()
    pdf.add_blank_page(width=300, height=300)
    stream = io.BytesIO()
    pdf.write(stream)
    try:
        parse_document("scan.pdf", stream.getvalue(), "before", "d3")
    except ValueError as exc:
        assert "OCR" in str(exc)
    else:
        raise AssertionError("Пустой скан ошибочно принят")


def test_openai_adapter_uses_structured_response_without_repeated_evidence(monkeypatch):
    from types import SimpleNamespace
    from backend.app.config import Settings
    from backend.app.provider import OpenAIProvider

    captured = {}

    class FakeResponses:
        def parse(self, **kwargs):
            captured.update(kwargs)
            return SimpleNamespace(status="completed", output_parsed=Matches(matches=[
                Match(before_id="f1", after_ids=["f2"], status="transferred",
                      explanation="Функция передана", recommendation="Проверить")]),
                usage=SimpleNamespace(input_tokens=80, output_tokens=25))

    monkeypatch.setattr("backend.app.provider.OpenAI", lambda **kwargs: SimpleNamespace(responses=FakeResponses()))
    provider = OpenAIProvider(Settings(_env_file=None, openai_api_key="test-placeholder"))
    old = {"id": "f1", "owner": "Отдел А", "action": "ведёт", "object": "реестр", "scope": "договоры", "evidence": [{"quote": "секретный фрагмент"}]}
    new = {**old, "id": "f2", "owner": "Отдел Б"}
    matches = provider.match([old], [new])
    assert matches.matches[0].status == "transferred"
    assert captured["store"] is False and captured["text_format"] is Matches
    assert "секретный фрагмент" not in str(captured["input"])
    assert provider.usage == {"input_tokens": 80, "output_tokens": 25}


def test_uncertain_is_not_published_as_a_loss():
    class UncertainProvider(BaselineProvider):
        def match(self, before, after):
            result = super().match(before, after)
            result.matches[0] = Match(before_id=before[0]["id"], after_ids=[], status="uncertain",
                                      explanation="Соответствие нельзя подтвердить", recommendation="Проверить вручную")
            return result

    result = analyze(fixture_documents("control"), UncertainProvider(), "openai")
    uncertain = next(f for f in result["findings"] if f["kind"] == "uncertain")
    assert uncertain["evidence_check"]["exact_quotes"]
    assert uncertain["search_scope"]["catalog_size"] == result["coverage"]["after_functions"]
    assert uncertain["search_scope"]["candidates"]
    assert all(e["side"] == "before" for e in uncertain["evidence"])


def test_full_ai_pipeline_contract_without_external_key(monkeypatch):
    import json
    from types import SimpleNamespace
    from backend.app.config import Settings
    from backend.app.provider import OpenAIProvider
    from backend.app.schemas import Clause, Extraction, UnitChanges, Risks

    baseline = BaselineProvider()

    class FakeResponses:
        def parse(self, **kwargs):
            schema = kwargs["text_format"]
            payload = json.loads(kwargs["input"][1]["content"])
            if schema is Extraction:
                output = baseline.extract([Clause.model_validate(c) for c in payload["target_clauses"]], payload["context"])
            elif schema is UnitChanges:
                output = baseline.units(payload["before"], payload["after"])
            elif schema is Matches:
                output = baseline.match(payload["before"], payload["after"])
            elif schema is Risks:
                output = baseline.risks(payload["focus"], payload["catalog"])
            else:
                raise AssertionError(schema)
            return SimpleNamespace(status="completed", output_parsed=output,
                                   usage=SimpleNamespace(input_tokens=1, output_tokens=1))

    monkeypatch.setattr("backend.app.provider.OpenAI", lambda **kwargs: SimpleNamespace(responses=FakeResponses()))
    provider = OpenAIProvider(Settings(_env_file=None, openai_api_key="test-placeholder"))
    result = analyze(fixture_documents("control"), provider, "openai")
    assert result["mode"] == "openai" and result["calls"] >= 4
    assert result["counts"]["not_found"] >= 1
    assert result["counts"]["duplication"] >= 1
    assert all(f["evidence_check"]["exact_quotes"] for f in result["findings"])
